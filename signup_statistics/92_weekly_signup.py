#!/usr/bin/env python3
"""
92. 주간 가입 보고 집계 — 조회용. BigQuery 와 파일 어느 것도 바꾸지 않는다

  정의는 README 「예비 정책 — 10/15 ~ 10/21 주 보고부터 적용」.
  SQL 만으로 못 하는 이유: 같은 회사 판정에 로그인 아이디가 필요한데 피드에 없다.
  관리자 export(cartaxbiz-usage_YYYY-MM-DD.xlsx)의 아이디 · 가입유형을 함께 쓴다.

  같은 회사 (연쇄로 묶는다)
    1. 사업자번호가 같다
    2. 정규화 회사명과 로그인 아이디가 둘 다 같다
    3. company_match_decisions.csv 에서 same 으로 판정했다
    판정 요청 = 정규화 회사명만 같고 사업자번호가 둘 다 없거나 한쪽에만 있다.
               판정 전에는 다른 회사로 센다. 이미 판정한 쌍은 다시 묻지 않는다.

  분류 (새 계정 가입 시점 기준)
    신규    같은 회사의 이전 계정이 없다
    중복    이전 계정 중 하나라도 가입 전날 유료(plus · premium)였다
    중복    이전 계정의 마지막 활동 후 30일 미만이다
    재가입  그 외
    마지막 활동 = 운행 · 관리자 로그인 · 가입 중 가장 늦은 것, 가입 전날 스냅샷 기준.

  개인정보: 아이디는 판정 요청 출력에만 보인다. 파일로 저장하지 않는다.
  담당자 · 연락처 열은 읽지 않는다.

  사용
    python3 -I signup_statistics/92_weekly_signup.py \
        --usage ../_inbox/cartaxbiz_usage/cartaxbiz-usage_2026-10-22.xlsx \
        --snap 2026-10-22 --week 2026-10-15

  --week 는 보고 주의 목요일. 그 주와 직전 주를 같이 낸다.
  BigQuery 읽기 2회(스냅샷 1개 + 이전 계정의 가입 전날 행). 수 MB.
"""
import argparse
import collections
import csv
import datetime as dt
import json
import os
import subprocess

import openpyxl

PROJECT = "carbiz-6f7fc"
# 판정에 안 걸린 테스트 계정 — README 「주간 가입 리포트 정의」 테스트 제외
MANUAL_TEST = {"X176", "H1424", "T1436"}
GAP_DAYS = 30
PAID = {"plus", "premium"}
HERE = os.path.dirname(os.path.abspath(__file__))
DECISIONS = os.path.join(HERE, "company_match_decisions.csv")
# 관리자 export 열 위치 (헤더 행으로 확인한다)
COL = {"가입일": 0, "회사명": 1, "코드": 2, "아이디": 4, "가입유형": 14}
TYPE_MAP = {"PC": "PC", "PC(데모)": "PC", "Mobile": "모바일", "AOS": "AOS", "iOS": "iOS"}


def bq(sql):
    out = subprocess.run(
        ["bq", "query", "--use_legacy_sql=false", "--format=json", "--max_rows=1000000"],
        input=sql, capture_output=True, text=True, check=True).stdout
    return json.loads(out[out.index("["):]) if "[" in out else []


def load_snapshot(snap):
    return bq(f"""
SELECT company_code AS c, company_name AS n,
       IFNULL(company_name_norm, '') AS nn, IFNULL(corporate_number_norm, '') AS corp,
       is_test_account AS t,
       FORMAT_DATETIME('%Y-%m-%dT%H:%M:%S', signup_date) AS s
FROM `{PROJECT}.signup_statistics.view_signup_statistics`
WHERE snapshot_date = '{snap}' AND signup_date IS NOT NULL""")


def load_usage(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    header = next(ws.iter_rows(max_row=1, values_only=True))
    for name, i in COL.items():
        if header[i] != name:
            raise SystemExit(f"export 열 위치가 바뀌었다: {i}번 열이 '{header[i]}' (기대 '{name}')")
    usage = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        usage[r[COL["코드"]]] = {
            "id": str(r[COL["아이디"]]).strip().lower() if r[COL["아이디"]] else "",
            "type": TYPE_MAP.get(r[COL["가입유형"]], "미상"),
        }
    return usage


def load_decisions():
    same, decided = [], set()
    with open(DECISIONS, encoding="utf-8") as f:
        for d in csv.DictReader(f):
            pair = frozenset((d["new_code"], d["prior_code"]))
            decided.add(pair)
            if d["answer"] == "same":
                same.append(tuple(pair))
    return same, decided


class Groups:
    def __init__(self, codes):
        self.p = {c: c for c in codes}

    def find(self, x):
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        if a in self.p and b in self.p:
            self.p[self.find(a)] = self.find(b)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--usage", required=True)
    ap.add_argument("--snap", required=True, help="계산에 쓸 스냅샷 날짜 — 보고에 함께 남긴다")
    ap.add_argument("--week", required=True, help="보고 주의 목요일 YYYY-MM-DD")
    a = ap.parse_args()

    this_start = dt.date.fromisoformat(a.week)
    if this_start.weekday() != 3:
        raise SystemExit("--week 는 목요일이어야 한다")
    weeks = [("지난주", this_start - dt.timedelta(days=7)), ("이번주", this_start)]

    usage = load_usage(a.usage)
    same, decided = load_decisions()
    rows = [r for r in load_snapshot(a.snap)
            if r["t"] in (False, "false") and r["c"] not in MANUAL_TEST]
    by_code = {r["c"]: r for r in rows}

    g = Groups(by_code)
    keyed = collections.defaultdict(list)
    for r in rows:
        if r["corp"]:
            keyed[("corp", r["corp"])].append(r["c"])
        lid = usage.get(r["c"], {}).get("id", "")
        if r["nn"] and lid:
            keyed[("name_id", r["nn"], lid)].append(r["c"])
    for codes in keyed.values():
        for c in codes[1:]:
            g.union(codes[0], c)
    for x, y in same:
        g.union(x, y)

    members = collections.defaultdict(list)
    by_name = collections.defaultdict(list)
    for r in rows:
        members[g.find(r["c"])].append(r)
        if r["nn"]:
            by_name[r["nn"]].append(r)

    def in_week(r, start):
        d = dt.date.fromisoformat(r["s"][:10])
        return start <= d <= start + dt.timedelta(days=6)

    targets = [r for r in rows if any(in_week(r, s) for _, s in weeks)]

    # 이전 계정의 가입 전날 행
    pairs, review = [], []
    for r in targets:
        prev_day = dt.date.fromisoformat(r["s"][:10]) - dt.timedelta(days=1)
        for p in members[g.find(r["c"])]:
            if p["s"] < r["s"]:
                pairs.append((r["c"], p["c"], prev_day.isoformat()))
        for p in by_name.get(r["nn"], []) if r["nn"] else []:
            if (p["s"] < r["s"] and g.find(p["c"]) != g.find(r["c"])
                    and not (p["corp"] and r["corp"])
                    and frozenset((r["c"], p["c"])) not in decided):
                review.append((r, p))

    prev = collections.defaultdict(list)
    if pairs:
        structs = ",".join(f"STRUCT('{n}' AS nc, '{p}' AS pc, DATE '{d}' AS pd)" for n, p, d in pairs)
        lo = min(d for _, _, d in pairs)
        for q in bq(f"""
SELECT x.nc, x.pc, IFNULL(v.plan_detail, '') AS pl,
       IFNULL(FORMAT_DATETIME('%Y-%m-%dT%H:%M:%S', v.last_trip_date_start), '') AS lt,
       IFNULL(FORMAT_DATETIME('%Y-%m-%dT%H:%M:%S', v.pc_last_login_date), '') AS ll,
       IFNULL(v.trip_count_total, 0) AS trips
FROM UNNEST([{structs}]) AS x
LEFT JOIN `{PROJECT}.signup_statistics.view_signup_statistics` AS v
  ON v.company_code = x.pc AND v.snapshot_date = x.pd
 AND v.snapshot_date BETWEEN '{lo}' AND '{a.snap}'"""):
            prev[q["nc"]].append(q)

    def classify(r):
        ps = prev.get(r["c"])
        if not ps:
            return "신규", []
        detail = []
        for q in ps:
            p = by_code[q["pc"]]
            acts = [p["s"]] + [v for v in (q["lt"], q["ll"]) if v and v < r["s"]]
            last = max(acts)
            gap = (dt.date.fromisoformat(r["s"][:10]) - dt.date.fromisoformat(last[:10])).days
            detail.append({"pc": q["pc"], "ps": p["s"][:10], "gap": gap,
                           "paid": q["pl"] in PAID, "trips": int(q["trips"])})
        detail.sort(key=lambda d: d["ps"], reverse=True)
        if any(d["paid"] for d in detail) or min(d["gap"] for d in detail) < GAP_DAYS:
            return "중복", detail
        return "재가입", detail

    result = {r["c"]: classify(r) for r in targets}

    print(f"스냅샷 {a.snap} · export {os.path.basename(a.usage)}\n")
    counts = {}
    for label, start in weeks:
        w = [r for r in targets if in_week(r, start)]
        c = collections.Counter(result[r["c"]][0] for r in w)
        t = collections.Counter(usage.get(r["c"], {}).get("type", "미상")
                                for r in w if result[r["c"]][0] == "신규")
        counts[label] = (c, t)

    def line(name, this, last):
        diff = this - last
        mark = "▲" if diff > 0 else "▼" if diff < 0 else "-"
        print(f"{name}\t{this}\t{diff:+d}건({mark})\t{last}")

    print("항목\t이번주\t변동\t지난주")
    (c1, t1), (c0, t0) = counts["이번주"], counts["지난주"]
    line("회원 가입", c1["신규"], c0["신규"])
    for k in ("PC", "모바일", "AOS", "iOS"):
        line(f"∟ {k}", t1[k], t0[k])
    line("∟ PC/모바일", t1["PC"] + t1["모바일"], t0["PC"] + t0["모바일"])
    line("∟ APP", t1["AOS"] + t1["iOS"], t0["AOS"] + t0["iOS"])
    line("재가입", c1["재가입"], c0["재가입"])
    print(f"(중복 — 보고 표에는 넣지 않는다)\t{c1['중복']}\t\t{c0['중복']}")
    if t1["미상"] or t0["미상"]:
        print(f"! 가입유형 미상: 이번주 {t1['미상']} · 지난주 {t0['미상']} — export 에 없는 계정")

    this_w = sorted((r for r in targets if in_week(r, this_start)), key=lambda r: r["s"])
    for kind in ("재가입", "중복"):
        print(f"\n== {kind} 목록 (이번주)")
        print("코드\t회사명\t이전 계정\t누적 운행\t이전 가입일\t마지막 이후 경과일")
        for r in this_w:
            k, detail = result[r["c"]]
            if k != kind:
                continue
            for d in detail:
                paid = " (유료)" if d["paid"] else ""
                print(f"{r['c']}\t{r['n']}\t{d['pc']}{paid}\t{d['trips']:,}\t{d['ps']}\t{d['gap']:,}일")

    print(f"\n== 판정 요청 {len(review)}건 — 답을 받으면 company_match_decisions.csv 에 적는다")
    for r, p in review:
        def side(x):
            corp = "있음" if x["corp"] else "없음"
            return f"{x['c']} {x['n']} · 아이디 {usage.get(x['c'], {}).get('id', '?')} · 사업자번호 {corp} · 가입 {x['s'][:10]}"
        print(f"  새 계정  {side(r)}\n  이전 계정 {side(p)}\n")


if __name__ == "__main__":
    main()
